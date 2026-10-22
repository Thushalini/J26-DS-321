"""Read an uploaded file into pages of plain text. Nothing is understood here, only read.

    pages = parse_document("Proposal.pdf", data)     # [Page(number=1, text="..."), ...]

PDF pages keep their real page numbers, so a fact can be traced to "Proposal.pdf p.4".
Word, text and Markdown files have no fixed pages, so they come back as one page with number None.

Known limit: a scanned PDF is a picture of text with no text layer. It cannot be read without
OCR, so it fails with a plain "no text found" message instead of silently giving nothing.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import PurePath
from typing import Optional

MAX_BYTES = 10 * 1024 * 1024          # 10 MB

# file ending -> kind of reader
SUPPORTED = {".pdf": "pdf", ".docx": "docx", ".txt": "text", ".md": "text", ".markdown": "text"}


class DocumentError(ValueError):
    """The file cannot be used. The message is safe to show to the person who uploaded it."""


@dataclass(frozen=True)
class Page:
    number: Optional[int]     # 1, 2, 3 ... for PDFs; None when the format has no fixed pages
    text: str


def kind_of(filename: str) -> str:
    """'pdf', 'docx' or 'text'. Raises DocumentError for anything else."""
    ending = PurePath(filename or "").suffix.lower()
    if ending not in SUPPORTED:
        allowed = ", ".join(sorted(SUPPORTED))
        raise DocumentError(f"This file type is not supported. Upload one of: {allowed}.")
    return SUPPORTED[ending]


def _pdf(data: bytes) -> list[Page]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentError("This PDF is password-protected. Upload a copy without a password.")
        return [Page(i, page.extract_text() or "") for i, page in enumerate(reader.pages, start=1)]
    except DocumentError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise DocumentError("This PDF could not be read. It may be damaged.") from exc


def _docx(data: bytes) -> list[Page]:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types for a bad file
        raise DocumentError("This Word file could not be read. Save it again as .docx and retry.") from exc
    # Each paragraph on its own line, with a blank line between, so headings such as
    # "References" stay on a line of their own for the chunker.
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return [Page(None, "\n\n".join(parts))]


def _text(data: bytes) -> list[Page]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return [Page(None, data.decode(encoding))]
        except UnicodeDecodeError:
            continue
    return [Page(None, data.decode("latin-1"))]


def parse_document(filename: str, data: bytes) -> list[Page]:
    """The file's pages. Raises DocumentError with a message meant for the uploader."""
    kind = kind_of(filename)
    if not data:
        raise DocumentError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise DocumentError(f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB.")
    pages = {"pdf": _pdf, "docx": _docx, "text": _text}[kind](data)
    if not any(p.text.strip() for p in pages):
        if kind == "pdf":
            raise DocumentError("No text found in this PDF. It may be a scanned image; "
                                "upload a version with selectable text.")
        raise DocumentError("No text found in this file.")
    return pages
